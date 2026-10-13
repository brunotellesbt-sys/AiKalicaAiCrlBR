"""Generate encounter/item indices and the DOCX from actual native gym exports."""
import argparse,json,re,struct,subprocess,tempfile
from pathlib import Path
from collections import Counter
from docx import Document
from audit_battle_database import compiled
ROOT=Path(__file__).resolve().parents[2]
NATURES=['Hardy','Lonely','Brave','Adamant','Naughty','Bold','Docile','Relaxed','Impish','Lax','Timid','Hasty','Serious','Jolly','Naive','Modest','Mild','Quiet','Bashful','Rash','Calm','Gentle','Sassy','Careful','Quirky']


def label(name,prefix='SPECIES_'):
    return name.removeprefix(prefix).replace('_',' ').title()


def generate(source,output,reports):
    source=Path(source);output=Path(output);reports=Path(reports);output.mkdir(parents=True,exist_ok=True)
    catalog=json.loads((ROOT/'tools/hoenn/catalog_metadata.json').read_text())['species']
    ecology=json.loads((source/'.journey-reproducible-water').read_text())['ecology']
    gen4=json.loads((source/'.journey-final-ecology').read_text())['generation4']
    species={int(k):label(v['name']) for k,v in catalog.items()}
    database=compiled(source)
    with tempfile.TemporaryDirectory(dir='/tmp') as temporary:
        p=Path(temporary);(p/'probe.c').write_text('#include "global.h"\n#include "pokemon.h"\n#include <stddef.h>\nconst u32 probe[]={offsetof(struct SpeciesInfo,abilities),sizeof(enum Ability)};\n')
        subprocess.run([str(ROOT/'.local/arm-gcc/usr/bin/arm-none-eabi-gcc'),'-S','-iquote',str(source/'include'),'-mabi=apcs-gnu',str(p/'probe.c'),'-o',str(p/'probe.s')],check=True)
        ability_offset,ability_bytes=map(int,re.findall(r'\.word\s+(\d+)',(p/'probe.s').read_text()))
    ability_names={int(v):label(n,'ABILITY_') for n,v in re.findall(r'\b(ABILITY_\w+)\s*=\s*(\d+)',(source/'include/constants/abilities.h').read_text())}
    rom=(source/'pokeemerald.gba').read_bytes();info=database['compiled_layout']['symbols']['gSpeciesInfo'];stride=database['compiled_layout']['abi'][0]
    def ability_name(mon):
        offset=info+mon['species']*stride+ability_offset+mon['ability_slot']*ability_bytes
        value=int.from_bytes(rom[offset:offset+ability_bytes],'little')
        return ability_names.get(value,str(value))+(' (Hidden)' if mon['ability_slot']==2 else '')
    items={int(v):label(n,'ITEM_') for n,v in re.findall(r'\b(ITEM_\w+)\s*=\s*(\d+)',(source/'include/constants/items.h').read_text())}
    moves={i:label(v['name'] or 'MOVE_'+str(i),'MOVE_') for i,v in database['moves'].items()}
    for cap,locations in [(9,ecology['locations_data']),(4,gen4['locations'])]:
        body=f'# Encontros por local — gerações 1–{cap}\n\n'
        body+='Cada família possui um local reservado. Os estágios encontrados acompanham a média das insígnias de Kanto e Hoenn: 0–2, básicos; 3–5, básicos e intermediários; 6–8, intermediários e finais. O nível fica entre a média da equipe menos 5 e mais 2, limitado a 1–100.\n\n'
        body+='Os membros de uma família listados abaixo não aparecem todos ao mesmo tempo. Formas, estágios e filtros aquáticos determinam quais podem aparecer naquele encontro. Lendários, míticos e Paradoxos estão no índice separado.\n\n'
        for location in locations:
            body+='## '+label(location['section'],'MAPSEC_')+'\n\n'
            body+='Mapas: '+', '.join(label(m,'MAP_') for m in location['maps'])+'.\n\n'
            body+='| Família | Membros elegíveis | Raridade da família |\n|---|---|---|\n'
            for family in location['families']:
                body+='| '+label(family['name'])+' | '+', '.join(label(s['name']) for s in family['species'])+' | '+str(family['rarity'])+' |\n'
            body+='\n'
        (output/f'encontros-gen1-{cap}.md').write_text(body.rstrip()+'\n')
    special=json.loads((source/'.journey-sanctuaries').read_text())
    captures=special['captures']+json.loads((source/'.journey-paradox-caves').read_text())['paradox_captures']
    body='# Lendários, míticos e Paradoxos\n\nAs capturas novas exigem as 16 insígnias, antes de qualquer Liga. Dialga, Palkia e Giratina exigem também concluir a missão Galactic no navio. A luta de Giratina nessa missão é obrigatória e não permite captura; a captura opcional ocorre posteriormente na caverna. Nas edições até a geração 4, espécies com número nacional acima de 493 ficam indisponíveis.\n\n'
    body+='| Pokémon | Número nacional | Caverna | Acesso/ilha de origem |\n|---|---|---|---|\n'
    sites={x['map']:x for x in special['sites']}
    for c in sorted(captures,key=lambda c:c['national_dex']):
        site=sites.get(c['map'],{})
        body+='| '+label(c['species'])+' | '+str(c['national_dex'])+' | '+c['map']+' | '+str(site.get('surface',c.get('parent','')))+'; '+str(site.get('access',''))+' |\n'
    body+='\nOs encontros originais ligados à história são preservados. Este índice enumera os novos santuários; não substitui as localizações originais mostradas pela AREA da Pokédex.\n'
    (output/'lendarios-miticos-paradoxos.md').write_text(body)
    supplies=json.loads((source/'.journey-evolution-supplies').read_text())
    body='# Itens de evolução e Mega Stones\n\nO fornecedor de cada Centro Pokémon abaixo vende as pedras obtidas em suas expedições. Mega Stones custam 25.000; itens de evolução mantêm o preço original, ou 6.000 quando originalmente não tinham preço de venda. A Mega Ring é entregue ao visitar um fornecedor depois de duas insígnias somadas entre as regiões.\n\n'
    body+='As evoluções por pedra do jogador continuam normais. Troca foi substituída por nível 36; troca com item, por nível 30 segurando o item. Os níveis mínimos 30/35/40 por força da evolução são uma regra adicional dos treinadores, não do jogador.\n\n'
    body+='| Local | Estoque |\n|---|---|\n'
    for shop in supplies['shops']:body+='| '+shop['map']+' | '+', '.join(label(n,'ITEM_') for n in shop['items'])+' |\n'
    body+='\n## Mega: espécie e pedra\n\n| Mega | Pedra |\n|---|---|\n'
    for row in supplies['mega_species_items']:body+='| '+label(row['species'])+' | '+label(row['item'],'ITEM_')+' |\n'
    body+='\nRayquaza usa Dragon Ascent, sem pedra. Os tutores nos laboratórios de Oak e Birch ensinam esse golpe; pedem confirmação antes de substituir o quarto golpe quando não há espaço. Esses tutores também ajustam as condições regionais de evolução para Local, Alola, Galar ou Hisui, sem mudar a região da história.\n\nAdamant Crystal, Lustrous Globe e Griseous Core são entregues apenas no final Galactic.\n\n## Alternativas por nível\n\n'
    for row in supplies['converted_evolutions']:body+='- '+label(row['target'])+': nível '+str(row['level'])+'.\n'
    body+='- Gholdengo: nível 36, sem a coleta de 999 moedas.\n- Urshifu Rapid Strike: nível 36 segurando Mystic Water; Single Strike: nível 36 sem esse item. Os scrolls também permanecem disponíveis.\n'
    (output/'itens-e-evolucoes.md').write_text(body)
    doc=Document();doc.add_heading('Pokémon Journey — equipes dos ginásios',0)
    doc.add_paragraph('Equipes extraídas diretamente das quatro ROMs, usando o gerador nativo de adversários. A ordem significa a posição do ginásio dentro da própria região, independentemente da outra região. As amostras cobrem as duas possibilidades de tamanho nos quatro primeiros ginásios e a equipe completa a partir do quinto.')
    doc.add_paragraph('1º: 2 adversários para 1–2 Pokémon do jogador, ou 3 para 3–6. 2º: 3 para 1–3, ou 4 para 4–6. 3º: 4 para 1–4, ou 5 para 5–6. 4º: 5 para 1–5, ou 6 para 6. 5º–8º: 6. Tate e Liza continuam em dupla; é necessário ter dois Pokémon aptos para uma batalha dupla.')
    doc.add_paragraph('EVs na ordem HP/Atk/Def/SpA/SpD/Spe; orçamento 510. Hard: IVs 31. Normal: IVs aleatórios; os valores extraídos são amostras, não IVs fixos do líder. Golpes de TM, tutor e breeding são permitidos cedo; level-up respeita o nível. Evoluções de treinadores respeitam a cadeia e os mínimos por item.')
    for name in ['Gen1-4-Normal','Gen1-4-Hard','Gen1-9-Normal','Gen1-9-Hard']:
        report=json.loads((reports/name/'gym-teams.json').read_text())
        doc.add_heading(name,1);doc.add_paragraph('SHA-256: '+report['rom_sha256'])
        grouped={}
        for row in report['checks']:grouped.setdefault(row['leader'],[]).append(row)
        for leader,rows in grouped.items():
            doc.add_heading(label(leader,'TRAINER_'),2)
            for row in rows:
                rank=row['rank'];count=len(row['mons'])
                doc.add_heading(f'{rank}º ginásio — equipe de {count}',3)
                table=doc.add_table(rows=1,cols=6);table.style='Light Shading Accent 1'
                for cell,text in zip(table.rows[0].cells,['Pokémon / Lv','Golpes','Nature','Item','EVs','Habilidade']):cell.text=text
                for mon in row['mons']:
                    cells=table.add_row().cells
                    values=[species[mon['species']]+' / '+str(mon['level']),', '.join(moves.get(m,'—') for m in mon['moves'] if m),NATURES[mon['nature']],items.get(mon['item'],'—'),'/'.join(map(str,mon['evs'])),ability_name(mon)]
                    for cell,text in zip(cells,values):cell.text=text
    doc.save(output/'equipes-ginasios-quatro-versoes.docx')
    (output/'gym-teams-native.json').write_text(json.dumps({name:json.loads((reports/name/'gym-teams.json').read_text()) for name in ['Gen1-4-Normal','Gen1-4-Hard','Gen1-9-Normal','Gen1-9-Hard']},indent=2)+'\n')
    print('Documentation and native gym DOCX generated:',output)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source',required=True);p.add_argument('--output',required=True);p.add_argument('--reports',required=True)
    a=p.parse_args();generate(a.source,a.output,a.reports)
